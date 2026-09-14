"""Stream standard OOXML workbooks; read them back with openpyxl."""
from pathlib import Path
import json
import zipfile
from xml.sax.saxutils import escape
import numpy as np
from openpyxl import load_workbook
from reproduce import sha, save


def write(path, sheets):
    path = Path(path)
    if path.exists():
        raise FileExistsError(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    def column(n):
        name = ""
        while n:
            n, remainder = divmod(n - 1, 26)
            name = chr(65 + remainder) + name
        return name
    columns = [column(i) for i in range(1, 40)]
    styles = f'''<styleSheet xmlns="{ns}"><numFmts count="2"><numFmt numFmtId="164" formatCode="0.0000"/><numFmt numFmtId="165" formatCode="0.00"/></numFmts><fonts count="2"><font><sz val="10"/><name val="Arial"/></font><font><b/><sz val="10"/><name val="Arial"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFD9EAF7"/></patternFill></fill></fills><borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="4"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="165" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>' + ''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, len(sheets) + 1)) + '</Types>')
        archive.writestr("_rels/.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        archive.writestr("xl/workbook.xml", f'<workbook xmlns="{ns}" xmlns:r="{rel}"><sheets>' + ''.join(f'<sheet name="{escape(s[0])}" sheetId="{i}" r:id="rId{i}"/>' for i, s in enumerate(sheets, 1)) + '</sheets></workbook>')
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + ''.join(f'<Relationship Id="rId{i}" Type="{rel}/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets) + 1)) + f'<Relationship Id="styles" Type="{rel}/styles" Target="styles.xml"/></Relationships>')
        archive.writestr("xl/styles.xml", styles)
        for sheet_id, (name, header, rows, raw) in enumerate(sheets, 1):
            with archive.open(f"xl/worksheets/sheet{sheet_id}.xml", "w") as stream:
                stream.write((f'<worksheet xmlns="{ns}"><sheetViews><sheetView workbookViewId="0"><pane xSplit="1" ySplit="1" topLeftCell="B2" activePane="bottomRight" state="frozen"/></sheetView></sheetViews><cols><col min="1" max="1" width="27" customWidth="1"/><col min="2" max="27" width="12" customWidth="1"/></cols><sheetData><row r="1" ht="36" customHeight="1">').encode())
                for j, value in enumerate(header):
                    if value is None:
                        continue
                    content = f'<is><t>{escape(value)}</t></is>' if isinstance(value, str) else f'<v>{value}</v>'
                    kind = ' t="inlineStr"' if isinstance(value, str) else ''
                    stream.write(f'<c r="{columns[j]}1" s="3"{kind}>{content}</c>'.encode("utf-8"))
                stream.write(b'</row>')
                for i, values in enumerate(rows, 2):
                    cells = [f'<row r="{i}">']
                    for j, value in enumerate(values):
                        if value is None or not np.isfinite(value):
                            continue
                        number = f"{float(value):.4f}" if j > 0 and not raw else f"{float(value):.17g}"
                        cells.append(f'<c r="{columns[j]}{i}" s="{2 if j == 0 else 1}"><v>{number}</v></c>')
                    cells.append('</row>')
                    stream.write(''.join(cells).encode())
                stream.write(b'</sheetData></worksheet>')


def field_rows(times, values):
    for t, row in zip(times, values):
        yield [float(t), *row]


def export(root, dest):
    dest.mkdir(parents=True, exist_ok=True)
    input_paths = [root / name for name in [
        "q1_complete_delivery/q1_delivery/output/q1_unrounded.npz",
        "q4_complete_delivery/v1/q123_closeout/output/q23_unified.npz",
        "q3_refinement_delivery/output/solution.npz",
        "q4_complete_delivery/v1/output/main.npz"]]
    input_hashes = {p.relative_to(root).as_posix(): sha(p) for p in input_paths}
    checkpoint = dest / "export_manifest.json"
    if checkpoint.exists():
        old = json.loads(checkpoint.read_text(encoding="utf-8"))
        assert old["inputs"] == input_hashes and old["builder_sha256"] == sha(Path(__file__))
        assert all(sha(dest / name) == digest for name, digest in old["outputs"].items())
        print("VERIFIED RESUME portable workbook export", flush=True)
        return
    q1 = np.load(root / "q1_complete_delivery/q1_delivery/output/q1_unrounded.npz")
    q2 = np.load(root / "q4_complete_delivery/v1/q123_closeout/output/q23_unified.npz")
    q3 = np.load(root / "q3_refinement_delivery/output/solution.npz")
    q4 = np.load(root / "q4_complete_delivery/v1/output/main.npz")
    radii = [j / 10 for j in range(21)]
    write(dest / "result1.xlsx", [
        ("温度", ["时间/s；径向距离/cm；温度/℃", *radii], field_rows(q1["time_s"][1:], q1["temperature_degC"][1:]), False),
        ("水分浓度", ["时间/s；径向距离/cm；干基含水率/(kg/kg)", *radii], field_rows(q1["time_s"][1:], q1["moisture_dry_basis"][1:]), False)])
    print("EXPORTED result1.xlsx", flush=True)
    for filename, mask in [("result2.xlsx", q2["time_s"] > 0),
                           ("result2_前3小时.xlsx", (q2["time_s"] > 0) & (q2["time_s"] <= 10800))]:
        write(dest / filename, [(name, ["时间/s；径向距离/cm", *radii],
               field_rows(q2["time_s"][mask], q2["profile_TC"][mask, :, ch]), False)
              for name, ch in [("温度", 0), ("水分浓度", 1)]])
        print("EXPORTED " + filename, flush=True)
    mask = q2["time_s"] > 0
    np.savez_compressed(dest / "result2_全程轨迹_206907x21.npz",
                        time_s=q2["time_s"][mask], radius_cm=q2["radius_cm"],
                        temperature_TC=np.round(q2["profile_TC"][mask, :, 0], 4),
                        moisture_C=np.round(q2["profile_TC"][mask, :, 1], 4))
    write(dest / "result3.xlsx", [("Sheet1", ["时间/s\\到药材中心的距离/cm", *radii],
                                 field_rows(q3["time_s"][1:], q3["profile_TC"][1:, :, 1]), True)])
    def q4_rows():
        for i in range(1, len(q4["time_s"])):
            yield [q4["time_s"][i], *q4["sample_TC"][i, :, 1], q4["surface_TC"][i, 1],
                   None, q4["radius_m"][i] * 100]
    write(dest / "result4.xlsx", [("Sheet1", ["时间/s\\到药材中心的距离/cm", *radii,
                                  "药材表面", None, "实际半径R(t) / cm"], q4_rows(), False)])
    print("EXPORTED result3.xlsx and result4.xlsx", flush=True)
    save(checkpoint, {"inputs": input_hashes, "builder_sha256": sha(Path(__file__)),
                      "outputs": {p.name: sha(p) for p in sorted(dest.iterdir()) if p.is_file()}})


def compare_workbook_to_arrays(path, arrays):
    """Read every exported numerical cell independently using openpyxl."""
    book = load_workbook(path, read_only=True, data_only=True)
    assert len(book.worksheets) == len(arrays), "Missing or extra workbook sheet"
    report = []
    for sheet, (times, values, rounded) in zip(book, arrays):
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        count = mismatches = 0
        for i, row in enumerate(rows):
            assert i < len(times), "Extra workbook row"
            assert abs(float(row[0]) - times[i]) < 1e-8, (sheet.title, i, "time")
            for j, expected in enumerate(values[i], 1):
                actual = row[j] if j < len(row) else None
                if expected is None or not np.isfinite(expected):
                    assert actual is None, (sheet.title, i, j, "outside-domain cell")
                else:
                    value = float(f"{float(expected):.4f}") if rounded else float(expected)
                    if actual is None or not np.isfinite(float(actual)) or abs(float(actual) - value) > (1e-10 if rounded else 1e-13):
                        mismatches += 1
                count += 1
        assert i + 1 == len(times), "Missing workbook rows"
        report.append({"sheet": sheet.title, "rows": len(times), "cells": count, "mismatches": mismatches})
    book.close()
    assert not any(r["mismatches"] for r in report), (str(path), report)
    return report
