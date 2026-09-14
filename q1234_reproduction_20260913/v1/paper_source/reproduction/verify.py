"""Compare regenerated numbers with paper/reference data; never feed solvers."""
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from reproduction.portable_workbooks import compare_workbook_to_arrays
from reproduce import sha


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def at(times, wanted):
    index = int(np.argmin(abs(np.asarray(times) - wanted)))
    assert abs(float(times[index]) - wanted) < 1e-6, (wanted, times[index])
    return index


def fmt(x, decimals=4):
    return "--" if x is None or not np.isfinite(x) else f"{float(x):.{decimals}f}"


def data(root):
    return [np.load(root / path) for path in [
        "q1_complete_delivery/q1_delivery/output/q1_unrounded.npz",
        "q4_complete_delivery/v1/q123_closeout/output/q23_unified.npz",
        "q3_refinement_delivery/output/solution.npz",
        "q4_complete_delivery/v1/output/main.npz"]]


def verify(package, root, out):
    q1, q2, q3, q4 = data(root)
    expected = load(package / "reference/paper_tables.json")["tables"]
    result = {"scope": "New PDE solves versus the saved paper display and submitted result values",
              "tables": {}, "claims": [], "full_outputs": {}}
    def claim(name, actual, expected_value, tolerance):
        ok = abs(actual - expected_value) <= tolerance
        result["claims"].append({"name": name, "actual": float(actual), "reference": expected_value,
                                 "absolute_tolerance": tolerance, "pass": bool(ok)})
    positions = [0, 5, 10, 15, 20]
    for name, table in expected.items():
        rows = []
        for old in table["rows"]:
            t = float(old[0]) * (1 if table["time_unit"] == "s" else 3600)
            if name.startswith("q1"):
                key = "temperature_degC" if name.endswith("temperature") else "moisture_dry_basis"
                values = q1[key][at(q1["time_s"], t), positions]
            elif name.startswith("q2"):
                channel = 0 if name.endswith("temperature") else 1
                values = q2["profile_TC"][at(q2["time_s"], t), positions, channel]
            elif name.startswith("q3"):
                values = q3["profile_TC"][at(q3["time_s"], t), positions, 1]
            else:
                i = at(q4["time_s"], t)
                values = np.r_[q4["sample_TC"][i, positions, 1], q4["surface_TC"][i, 1]]
            rows.append([str(int(t)) if table["time_unit"] == "s" else fmt(t / 3600), *[fmt(v) for v in values]])
        mismatch = [{"row": i + 1, "column": j + 1, "actual": a, "reference": b}
                    for i, (actual, target) in enumerate(zip(rows, table["rows"]))
                    for j, (a, b) in enumerate(zip(actual, target)) if a != b]
        result["tables"][name] = {"pass": not mismatch, "rows": rows,
                                  "numeric_cells": sum(v != "--" for r in rows for v in r), "mismatches": mismatch}
    q3j = load(root / "q3_refinement_delivery/output/solution.json")["event"]
    q4j = load(root / "q4_complete_delivery/v1/output/end_event.json")
    claim("Q3 critical seconds", q3j["critical_s"], 206906.389932328, 0.0005)
    claim("Q3 execution hours", q3j["execution_h"], 57.4741, 1e-10)
    claim("Q3 execution maximum C", q3j["execution_max"], 0.149999892208, 1e-9)
    claim("Q4 critical seconds", q4j["critical_s"], 183931.211715489, 0.0005)
    claim("Q4 execution hours", q4j["execution_h"], 51.0921, 1e-10)
    assert q3j["execution_max"] < .15 and q4j["execution_max_C"] < .15
    fixed = load(root / "q4_complete_delivery/v1/validation/fixed80.json")["event"]["critical_h"]
    claim("Q4 same-property fixed-radius critical hours", fixed, 129.8484, 0.00005)
    claim("Q4 time reduction percent", 100 * (1 - q4j["critical_h"] / fixed), 60.6526, 0.00005)
    # Verify every saved Q2 value, without equating binary files with numerical equality.
    old = np.load(package / "results/result2_全程轨迹_206907x21.npz")
    mask = q2["time_s"] > 0
    np.testing.assert_allclose(q2["time_s"][mask], old["time_s"], rtol=0, atol=1e-8)
    for key, ch in [("temperature_TC", 0), ("moisture_C", 1)]:
        actual = q2["profile_TC"][mask, :, ch]
        target = old[key]
        assert actual.shape == target.shape
        delta = abs(actual - target)
        rounding_difference = np.round(actual, 4) != target
        ii = np.argwhere(rounding_difference)
        result["full_outputs"]["q2_" + key] = {
            "cells": int(target.size), "max_abs_unrounded_to_submitted_rounded": float(delta.max()),
            "four_decimal_different_cells": int(rounding_difference.sum()),
            "first_differences": [{"time_s": float(old["time_s"][i]), "radius_cm": float(old["radius_cm"][j]),
                                   "new_unrounded": float(actual[i, j]), "submitted": float(target[i, j])} for i, j in ii[:12]],
            "pass": bool(np.all(delta <= 0.00005002)),
            "tolerance_reason": "Submitted values are rounded to 4 decimals; half a display unit plus 2e-8 numerical tolerance"}
    # Compare all workbook values from the other three questions at their displayed precision.
    for case, times, values in [(1, q1["time_s"][1:], [q1["temperature_degC"][1:], q1["moisture_dry_basis"][1:]]),
                                (3, q3["time_s"][1:], [q3["profile_TC"][1:, :, 1]]),
                                (4, q4["time_s"][1:], [np.c_[q4["sample_TC"][1:, :, 1], q4["surface_TC"][1:, 1]]])]:
        workbook = load_workbook(package / f"results/result{case}.xlsx", read_only=True, data_only=True)
        checked = differences = 0
        largest = 0.0
        for sheet, actual in zip(workbook, values):
            rows = sheet.iter_rows(values_only=True)
            next(rows)
            for i, row in enumerate(rows):
                assert i < len(times) and abs(row[0] - times[i]) < 1e-7
                for j, x in enumerate(actual[i], 1):
                    old_value = row[j] if j < len(row) else None
                    if not np.isfinite(x):
                        assert old_value is None
                    else:
                        assert old_value is not None
                        largest = max(largest, abs(float(x) - old_value))
                        differences += fmt(x) != fmt(old_value)
                    checked += 1
            assert i + 1 == len(times)
        workbook.close()
        result["full_outputs"][f"q{case}_workbook"] = {"cells": checked, "four_decimal_different_cells": int(differences),
            "max_abs_difference": largest, "pass": largest <= (2e-8 if case == 3 else 0.00005002)}
    # Recompute the additional table of whole-cylinder averages.
    geometry = root / "q4_complete_delivery/v1/validation/geometry"
    result["geometry_table"] = []
    for case, pair, t, means, percentage in [
        ("q1", ["q1_80x128", "q1_80x0"], 1800, [2.274929, 2.293532], 7.2533),
        ("q2", ["q23_80x256_integral", "q23_80x0_integral"], 10800, [1.327482, 1.382492], 4.7118),
        ("q4", ["q4_80x128_integral", "q4_80x0_integral"], 21600, [1.024545, 1.065198], 2.7379)]:
        actual = []
        for stem in pair:
            z = np.load(geometry / (stem + ".npz"))
            field = z["snapshots"][at(z["snapshot_times_s"], t), ..., 1]
            actual.append(float(field.ravel() @ z["weights"] / z["weights"].sum()))
        pct = 100 * ((2.55 - actual[0]) / (2.55 - actual[1]) - 1)
        claim(case + " whole 2D mean", actual[0], means[0], 0.0000005)
        claim(case + " whole 1D mean", actual[1], means[1], 0.0000005)
        claim(case + " loss increase percent", pct, percentage, 0.00005)
        result["geometry_table"].append({"case": case, "time_s": t, "means_2d_1d": actual, "loss_increase_percent": pct})
    # Every displayed six-decimal environmental root and root difference.
    envdir = root / "q1234_overall_review_delivery/v1/evidence/environment"
    result["environment_table"] = []
    scenarios = ["baseline", "temperature_minus_1K", "temperature_plus_1K", "boundary_minus_0p005", "boundary_plus_0p005"]
    targets = {"q3": ([57.473996, 59.341275, 55.688357, 57.382748, 57.584113], [0, 1.867279, -1.785639, -0.091249, 0.110117]),
               "q4": ([51.092003, 52.729346, 49.525343, 50.993818, 51.218753], [0, 1.637343, -1.566660, -0.098185, 0.126750])}
    for case in ("q3", "q4"):
        baseline = load(envdir / f"{case}_baseline_n40.json")["critical_s"]
        for i, scenario in enumerate(scenarios):
            value = load(envdir / f"{case}_{scenario}_n40.json")["critical_s"]
            delta = (value - baseline) / 3600
            claim(f"{case} {scenario} root hours", value / 3600, targets[case][0][i], 0.0000005)
            claim(f"{case} {scenario} change hours", delta, targets[case][1][i], 0.0000005)
            result["environment_table"].append({"case": case, "scenario": scenario, "root_h": value / 3600, "delta_h": delta})
    sensitivity = load(root / "q2_refinement_delivery/reproduced_sensitivity/sensitivity.json")
    for name, center, surface in [("hm_0.8", 1.8581, 1.1775), ("hm_1.2", 1.6921, 0.8737)]:
        claim("Q2 " + name + " center at 3h", sensitivity[name]["C_center_end"], center, 0.00005)
        claim("Q2 " + name + " surface at 3h", sensitivity[name]["C_surface_end"], surface, 0.00005)
    endpoint = load(root / "q1234_overall_review_delivery/v1/reviews/geometry/endpoint_audit.json")
    previous_endpoint = load(package / "support/evidence/q4_endpoint_20260912/endpoint_audit.json")
    result["geometry_endpoint_reproduction"] = endpoint["cases"]
    for case in ["q23", "q4"]:
        claim(case + " coarse 2D Cmax at the stated time",
              endpoint["cases"][case]["direct_FV_Cmax_at_execution"],
              previous_endpoint["cases"][case]["direct_FV_Cmax_at_execution"], 1e-8)
    assert endpoint["cases"]["q23"]["all_nodes_strictly_pass"]
    assert not endpoint["cases"]["q4"]["all_nodes_strictly_pass"]
    result["paper_numeric_cells"] = sum(t["numeric_cells"] for t in result["tables"].values())
    result["pass"] = all(t["pass"] for t in result["tables"].values()) and all(c["pass"] for c in result["claims"]) and all(c["pass"] for c in result["full_outputs"].values())
    return result


def verify_exports(root, out):
    q1, q2, q3, q4 = data(root)
    report = {}
    specs = {
        "result1.xlsx": [(q1["time_s"][1:], q1[key][1:], True) for key in ["temperature_degC", "moisture_dry_basis"]],
        "result2.xlsx": [(q2["time_s"][1:], q2["profile_TC"][1:, :, c], True) for c in [0, 1]],
        "result2_前3小时.xlsx": [(q2["time_s"][1:10801], q2["profile_TC"][1:10801, :, c], True) for c in [0, 1]],
        "result3.xlsx": [(q3["time_s"][1:], q3["profile_TC"][1:, :, 1], False)],
        "result4.xlsx": [(q4["time_s"][1:], np.c_[q4["sample_TC"][1:, :, 1], q4["surface_TC"][1:, 1],
                                    np.full(len(q4["time_s"]) - 1, np.nan), q4["radius_m"][1:] * 100], True)]}
    for filename, arrays in specs.items():
        report[filename] = compare_workbook_to_arrays(out / "results" / filename, arrays)
        print("READ-BACK PASS " + filename, flush=True)
    return {"pass": True, "workbooks": report, "total_cells": sum(r["cells"] for rows in report.values() for r in rows),
            "export_manifest_sha256": sha(out / "results/export_manifest.json")}
